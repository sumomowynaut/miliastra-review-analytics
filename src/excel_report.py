"""
Excel 报告生成 —— 7 个工作表

自动校验可机器检查的部分（能否打开 / 表数量 / 非空 / 乱码）；
完整交付清单见 methodology/06-delivery-checklist.md。
"""

import re
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# 样式
FONT = "Microsoft YaHei"
H_FILL = PatternFill("solid", fgColor="1F3864")
H_FONT = Font(name=FONT, size=11, bold=True, color="FFFFFF")
T_FONT = Font(name=FONT, size=10)
BOLD = Font(name=FONT, size=10, bold=True)
TITLE = Font(name=FONT, size=14, bold=True, color="1F3864")
WARN = Font(name=FONT, size=10, bold=True, color="C00000")
GOOD = Font(name=FONT, size=10, color="1F7A3D")
BAD = Font(name=FONT, size=10, color="C00000")
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center")


def _header(ws, row, headers, widths=None):
    for i, h in enumerate(headers, 1):
        c = ws.cell(row=row, column=i, value=h)
        c.fill, c.font, c.border = H_FILL, H_FONT, BORDER
        c.alignment = CENTER
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w


def _row(ws, row, values, zebra=False):
    for i, v in enumerate(values, 1):
        c = ws.cell(row=row, column=i, value=v)
        c.font, c.border = T_FONT, BORDER
        c.alignment = WRAP
        if zebra:
            c.fill = PatternFill("solid", fgColor="F7F9FC")


def build_excel(analysis, details, out_path, level_id="", title="",
                meta=None, topics_coverage=None, updates=None):
    wb = Workbook()

    # ---------------------------------------------------- 1. 概览
    ws = wb.active
    ws.title = "概览"
    r = analysis["meta"]
    rates = r["rates"]
    ts = r["template_stats"]

    ws["A1"] = f"千星奇域 {level_id} 数据报告"
    ws["A1"].font = TITLE
    ws.merge_cells("A1:I1")
    ws["A2"] = f"{title}  |  生成时间 {datetime.now():%Y-%m-%d %H:%M}"
    ws["A2"].font = Font(name=FONT, size=9, color="808080")
    ws.merge_cells("A2:I2")

    row = 4
    _header(ws, row, ["指标", "表观（含模板）", "真实（已剥离）", "差异", "说明"],
            [26, 18, 18, 14, 46])
    row += 1
    ap, rl = rates["apparent"], rates["real"]
    for label, ak, rk, note in [
        ("总评论数", "total", "total", "表观含系统模板语"),
        ("好评率(%)", "good_rate", "good_rate", "表观口径与官方公布值对齐"),
        ("差评率(%)", "bad_rate", "bad_rate", "★ 真实差评率才是体验结论依据"),
    ]:
        _row(ws, row, [label, ap[ak], rl[rk],
                       (round(rl[rk] - ap[ak], 2) if isinstance(ap[ak], (int, float))
                        and isinstance(rl[rk], (int, float)) else ""), note])
        if label == "差评率(%)":
            ws.cell(row=row, column=3).font = BAD
            ws.cell(row=row, column=2).font = GOOD
        row += 1

    row += 1
    ws.cell(row=row, column=1, value="模板语剥离").font = BOLD
    row += 1
    _header(ws, row, ["模板文案", "出现次数", "占模板比例", "", "备注"], [26, 18, 18, 14, 46])
    row += 1
    for t in ts.get("top_templates", []):
        _row(ws, row, [t["text"], t["count"], round(t["pct_of_templates"], 1), "",
                       "自动探测"])
        row += 1
    _row(ws, row, ["合计", ts["template_count"],
                   round(ts["template_pct"], 1), "",
                   f"真实评论 {ts['real_count']} 条 / {ts['real_unique']} 种唯一文本"])
    ws.cell(row=row, column=1).font = BOLD
    row += 2

    ws.cell(row=row, column=1, value="重要提示").font = WARN
    row += 1
    # ★ 这里曾写死案例数字（官方公布的 79.2%），换一个奇域就会输出错误结论。
    #   有官方口径就用官方的，没有就用本报告算出的表观口径。
    official_rate = _fmt_rate((meta or {}).get("goodRate"))
    apparent_label = (f"官方公布的 {official_rate}" if official_rate
                      else f"表观好评率 {ap['good_rate']:.1f}%")
    for line in [
        "· 表观口径（含系统模板）用于与官方公布的好评率对齐，两者一致可证明数据完整。",
        "· 真实口径（已剥离模板）才反映玩家实际评价，差异 "
        f"{rates['delta_bad_rate']:.1f} 个百分点。",
        "· 所有分布、聚类、趋势结论均基于真实评论。",
        f"· {apparent_label} 是含模板口径，不宜直接用于判断真实体验。",
    ]:
        c = ws.cell(row=row, column=1, value=line)
        c.font = T_FONT
        c.alignment = Alignment(wrap_text=True)
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
        row += 1

    # ---------------------------------------------------- 2. 模板剥离分析
    ws2 = wb.create_sheet("模板剥离分析")
    ws2["A1"] = "模板语探测与剥离明细"
    ws2["A1"].font = TITLE
    ws2.merge_cells("A1:F1")
    _header(ws2, 3, ["模板文案", "次数", "占模板比", "探测来源", "归一化后", "样本原句"],
            [40, 10, 12, 12, 30, 40])
    row = 4
    for t in (ts.get("templates") or ts.get("top_templates", [])):
        _row(ws2, row, [t.get("sample") or t["text"], t["count"],
                        round(t["pct_of_templates"], 1), "自动探测",
                        t["text"], t.get("sample") or ""])
        row += 1
    user_counts = ts.get("user_added_counts", {})
    user_samples = ts.get("user_added_samples", {})
    for t in ts.get("user_added", []):
        shown = user_samples.get(t, t)
        n = user_counts.get(t, 0)
        pct = n / ts["template_count"] * 100 if ts["template_count"] else 0
        _row(ws2, row, [shown, n, round(pct, 1), "用户补充", t, shown])
        row += 1
    row += 1
    _row(ws2, row, ["剥离后验证", "", "", "", "", ""])
    ws2.cell(row=row, column=1).font = BOLD
    row += 1
    _row(ws2, row, ["真实评论数", ts["real_count"], "", "", "", ""])
    _row(ws2, row + 1, ["唯一文本数", ts["real_unique"], "", "", "真实评论应几乎无重复", ""])
    _row(ws2, row + 2, ["重复率(%)", round(ts["real_dup_rate"], 1), "",
                        "", "偏高说明仍有模板变体未探测到", ""])

    # ---------------------------------------------------- 3. 评论明细
    ws3 = wb.create_sheet("评论明细")
    ws3["A1"] = "评论明细（含模板标记）"
    ws3["A1"].font = TITLE
    ws3.merge_cells("A1:N1")
    headers = ["楼层", "回复ID", "日期", "用户", "UID", "IP属地", "点赞",
               "好评", "模板", "内容", "小时", "星期", "字数", "跟帖数"]
    _header(ws3, 3, headers,
            [8, 22, 12, 16, 14, 10, 8, 8, 8, 60, 8, 8, 8, 8])
    row = 4
    for i, d in enumerate(sorted(details, key=lambda x: (x.get("date", ""), -x.get("like", 0)))):
        dt = str(d.get("date", ""))
        hour, dow = "", ""
        try:
            parsed = datetime.strptime(dt[:19], "%Y-%m-%d %H:%M:%S")
            hour = f"{parsed.hour:02d}"
            dow = "周" + "一二三四五六日"[parsed.weekday()]
        except ValueError:
            m = re.search(r"(\d{2}):(\d{2})", dt)
            if m:
                hour = m.group(1)
        _row(ws3, row, [
            d.get("floor_id", ""), d.get("reply_id", ""), dt[:16],
            d.get("user", ""), d.get("uid", ""), d.get("ip", ""),
            d.get("like", 0),
            "✓" if d.get("is_recommend") else "",
            "是" if d.get("is_template") else "",
            d.get("content", ""), hour, dow,
            len(d.get("content", "")), d.get("reply_count", 0),
        ], zebra=(i % 2 == 1))
        row += 1
    ws3.freeze_panes = "A4"
    ws3.auto_filter.ref = f"A3:N{row-1}"

    # ---------------------------------------------------- 4. 月度趋势
    ws4 = wb.create_sheet("月度趋势")
    ws4["A1"] = "月度趋势（双口径，绝对量与占比并排）"
    ws4["A1"].font = TITLE
    ws4.merge_cells("A1:H1")
    _header(ws4, 3, ["月份", "真实总数", "真实差评数", "真实差评占比(%)",
                     "表观总数", "表观差评占比(%)", "样本提示", ""],
            [12, 12, 14, 18, 12, 18, 16, 8])
    row = 4
    ap_map = {m["month"]: m for m in analysis.get("monthly_apparent", [])}
    for m in analysis.get("monthly_real", []):
        am = ap_map.get(m["month"], {})
        _row(ws4, row, [m["month"], m["total"], m["bad"], round(m["bad_pct"], 1),
                        am.get("total", ""), round(am.get("bad_pct", 0), 1),
                        m.get("sample_note", "") or "—", ""])
        row += 1
    row += 1
    c = ws4.cell(row=row, column=1,
                 value="⚠️ 占比会被分母稀释 —— 差评绝对量与占比必须并排看。样本量 < 30 的分组不可用于结论。")
    c.font = WARN
    ws4.merge_cells(start_row=row, start_column=1, end_row=row, end_column=7)

    # ---------------------------------------------------- 5. 主题统计
    ws5 = wb.create_sheet("主题统计")
    ws5["A1"] = "痛点聚类（仅基于真实差评）"
    ws5["A1"].font = TITLE
    ws5.merge_cells("A1:F1")
    cov = analysis.get("coverage", {})
    ws5["A2"] = (f"差评总数 {cov.get('negative_total', 0)} | "
                 f"词典覆盖率 {cov.get('coverage', 0):.1f}% {cov.get('coverage_note', '')}")
    ws5["A2"].font = T_FONT
    _header(ws5, 4, ["痛点类别", "条数", "占差评比(%)", "代表评论1", "代表评论2", "代表评论3"],
            [18, 10, 14, 34, 34, 34])
    row = 5
    for t in analysis.get("topics", []):
        s = t.get("samples", [])
        _row(ws5, row, [
            t["category"], t["count"], round(t["pct_of_negative"], 1),
            s[0]["text"] if len(s) > 0 else "",
            s[1]["text"] if len(s) > 1 else "",
            s[2]["text"] if len(s) > 2 else "",
        ])
        row += 1
    if cov.get("unmatched"):
        row += 1
        c = ws5.cell(row=row, column=1,
                     value=f"未命中 {cov['unmatched']} 条（人工复核样例）")
        c.font = WARN
        row += 1
        for u in cov.get("unmatched_samples", []):
            _row(ws5, row, ["", "", "", u["date"], u["text"], ""])
            row += 1

    # ---------------------------------------------------- 6. 地域分布
    ws6 = wb.create_sheet("地域分布")
    ws6["A1"] = "IP 属地分布（真实评论）"
    ws6["A1"].font = TITLE
    _header(ws6, 3, ["IP属地", "评论数", "占比(%)", "", "", ""], [16, 12, 12, 8, 8, 8])
    row = 4
    ips = analysis.get("ip_distribution", [])
    tot = sum(x[1] for x in ips) or 1
    for name, cnt in ips:
        _row(ws6, row, [name, cnt, round(cnt / tot * 100, 1), "", "", ""])
        row += 1

    # ---------------------------------------------------- 7. 方法论与更新日志
    ws7 = wb.create_sheet("方法论与日志")
    ws7["A1"] = "分析方法与记录"
    ws7["A1"].font = TITLE
    ws7.merge_cells("A1:D1")
    _header(ws7, 3, ["类别", "条目", "说明", "来源"], [16, 34, 52, 22])
    row = 4
    for cat, item, desc, src in [
        ("数据完整性", "双趟排序合并", "FLOOR_DESC + HOT 各跑一趟合并；单趟会漏约 10%", "methodology/01"),
        ("数据完整性", "每页重试 3 次", "失败不 break，单页失败不能拖垮整趟", "methodology/01"),
        ("数据完整性", "点赞数取 max", "两趟数据不一致时取最大值，不覆盖", "methodology/01"),
        ("数据完整性", "完整性七项校验", "条数/回复ID/跨度/降级/趟产出/楼层完整/好评率对齐", "methodology/01"),
        ("模板剥离", "自动探测", "高频完全重复 = 模板；不硬编码平台文案", "methodology/02"),
        ("模板剥离", "双口径并报", "表观对齐官方，真实用于结论", "methodology/02"),
        ("痛点聚类", "词典召回 + 人工裁决", "报告未命中率；未命中率 >30% 需修词典", "methodology/03"),
        ("痛点聚类", "否定语义剔除", "「不卡」「不好玩」需从差评桶剔除", "methodology/03"),
        ("趋势分析", "绝对量 + 占比并排", "占比会被分母稀释", "methodology/04"),
        ("趋势分析", "样本 < 30 不下结论", "分组样本量不足需标注", "methodology/04"),
        ("结论纪律", "先验证再下结论", "每个结论写明推翻条件", "methodology/05"),
        ("结论纪律", "三类陈述分开标注", "官方事实 / 玩家陈述 / 我的推断", "methodology/05"),
    ]:
        _row(ws7, row, [cat, item, desc, src])
        row += 1
    if updates:
        row += 1
        c = ws7.cell(row=row, column=1, value="更新日志")
        c.font = BOLD
        row += 1
        _header(ws7, row, ["日期", "变更", "", ""], [16, 34, 52, 22])
        row += 1
        for u in updates:
            _row(ws7, row, [u.get("date", ""), u.get("change", ""), "", ""])
            row += 1

    for s in wb.worksheets:
        s.sheet_view.showGridLines = False
    wb.save(out_path)
    return out_path


def verify_excel(path, expect_sheets=7):
    """交付前校验 —— 实际重新打开检查"""
    from openpyxl import load_workbook
    wb = load_workbook(path)
    try:
        issues = []
        if len(wb.sheetnames) != expect_sheets:
            issues.append(f"工作表数量 {len(wb.sheetnames)} != {expect_sheets}")
        for ws in wb.worksheets:
            if ws.max_row < 2:
                issues.append(f"{ws.title} 为空")
        # 中文乱码抽检
        ws = wb.worksheets[0]
        if ws.max_row >= 2 and "�" in str(ws.cell(row=1, column=1).value or ""):
            issues.append("概览表头疑似乱码")
        return issues
    finally:
        # Windows 下不关句柄会锁住文件，导致后续删除/覆盖失败
        wb.close()


def _fmt_rate(value):
    """把官方好评率统一成「79.2%」形式，避免输出 79.177545691906"""
    s = str(value or "").strip()
    if not s:
        return ""
    if s.endswith("%"):
        return s
    try:
        return f"{float(s):.1f}%"
    except (TypeError, ValueError):
        return s
